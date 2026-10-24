"""
document_manager.py
-------------------
Fase 7.8.1 — Hotfix de reingestión.

Permite reprocesar el MISMO archivo (mismo document_id + mismo hash)
cuando la ingestión previa fue realizada con una versión anterior del parser.

Una vez procesado con PARSER_VERSION = "7.8",
el duplicado vuelve a bloquearse normalmente.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from datetime import datetime
from pathlib import Path

from .tabular_ingestion import (
    extract_csv_rows,
    extract_excel_rows,
    extract_pdf_tables,
    rows_to_chunks,
)


PARSER_VERSION = "7.8"

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".txt",
    ".csv",
    ".xlsx",
    ".xls",
}


def _jsonl_load(path: Path) -> list[dict]:
    if not path.exists():
        return []

    rows = []

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:
        for line in file:
            line = line.strip()

            if line:
                rows.append(
                    json.loads(line)
                )

    return rows


def _jsonl_write(
    path: Path,
    rows: list[dict],
):
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with path.open(
        "w",
        encoding="utf-8",
    ) as file:
        for row in rows:
            file.write(
                json.dumps(
                    row,
                    ensure_ascii=False,
                )
                + "\n"
            )


def _sha256(data: bytes) -> str:
    return hashlib.sha256(
        data
    ).hexdigest()


def _clean_text(text: str) -> str:
    text = str(text).replace(
        "\xa0",
        " ",
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


def _chunk_text(
    text: str,
    chunk_size: int,
    overlap: int,
) -> list[str]:

    text = _clean_text(
        text
    )

    if not text:
        return []

    chunk_size = max(
        int(chunk_size),
        200,
    )

    overlap = max(
        0,
        min(
            int(overlap),
            chunk_size - 1,
        ),
    )

    chunks = []

    start = 0
    length = len(text)

    while start < length:
        end = min(
            start + chunk_size,
            length,
        )

        chunk = text[
            start:end
        ].strip()

        if chunk:
            chunks.append(
                chunk
            )

        if end >= length:
            break

        start = end - overlap

    return chunks


def _extract_pdf_text(
    data: bytes,
) -> list[tuple[str, str]]:

    from pypdf import PdfReader

    reader = PdfReader(
        io.BytesIO(
            data
        )
    )

    pages = []

    for idx, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = (
            page.extract_text()
            or ""
        )

        pages.append(
            (
                f"Página {idx}",
                _clean_text(
                    text
                ),
            )
        )

    return pages


def _extract_docx_text(
    data: bytes,
) -> str:

    from docx import Document

    document = Document(
        io.BytesIO(
            data
        )
    )

    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(
                text
            )

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(values):
                parts.append(
                    " | ".join(
                        values
                    )
                )

    return _clean_text(
        "\n".join(
            parts
        )
    )


def _extract_txt(
    data: bytes,
) -> str:

    for encoding in (
        "utf-8-sig",
        "utf-8",
        "latin-1",
    ):
        try:
            return _clean_text(
                data.decode(
                    encoding
                )
            )

        except Exception:
            pass

    raise RuntimeError(
        "No fue posible decodificar TXT."
    )


def _build_generic_chunks(
    sections: list[tuple[str, str]],
    *,
    document_id: str,
    title: str,
    source_file: str,
    document_type: str,
    source_status: str,
    access_level: str,
    document_date: str,
    version: str,
    file_hash: str,
    ingested_at: str,
    chunk_size: int,
    overlap: int,
) -> list[dict]:

    result = []
    chunk_counter = 0

    for section, text in sections:

        for chunk_text in _chunk_text(
            text,
            chunk_size,
            overlap,
        ):
            chunk_counter += 1

            result.append(
                {
                    "chunk_id": (
                        f"{document_id}-C"
                        f"{chunk_counter:04d}"
                    ),
                    "document_id": document_id,
                    "title": title,
                    "source_file": source_file,
                    "document_type": document_type,
                    "section": section,
                    "chunk_index": chunk_counter,
                    "text": chunk_text,
                    "source_status": source_status,
                    "access_level": access_level,
                    "document_date": document_date,
                    "version": version,
                    "file_hash": file_hash,
                    "ingested_at": ingested_at,
                    "content_mode": "generic_text",
                    "parser_version": PARSER_VERSION,
                }
            )

    return result


def ingest_uploaded_file(
    *,
    uploaded_file,
    data_dir,
    document_id: str,
    title: str,
    document_type: str,
    source_status: str,
    access_level: str,
    document_date: str,
    version: str,
    chunk_size: int = 1200,
    overlap: int = 200,
):
    """
    Reglas de reingestión:

    - mismo ID + mismo hash + parser_version actual:
        BLOQUEAR (duplicado real)

    - mismo ID + mismo hash + parser_version antiguo o inexistente:
        REPROCESAR automáticamente

    - mismo ID + hash distinto:
        REEMPLAZAR documento anterior
    """

    data_dir = Path(
        data_dir
    )

    private_dir = (
        data_dir
        / "private_ingested"
    )

    files_dir = (
        private_dir
        / "files"
    )

    private_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    files_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    registry_path = (
        private_dir
        / "document_registry.jsonl"
    )

    chunks_path = (
        private_dir
        / "knowledge_chunks.jsonl"
    )

    filename = str(
        uploaded_file.name
    )

    extension = (
        Path(
            filename
        )
        .suffix
        .lower()
    )

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Formato no soportado: {extension}"
        )

    file_bytes = uploaded_file.getvalue()

    if not file_bytes:
        raise ValueError(
            "El archivo está vacío."
        )

    document_id = (
        document_id.strip()
    )

    title = (
        title.strip()
    )

    if not document_id:
        raise ValueError(
            "document_id es obligatorio."
        )

    if not title:
        raise ValueError(
            "title es obligatorio."
        )

    file_hash = _sha256(
        file_bytes
    )

    ingested_at = (
        datetime.now()
        .isoformat(
            timespec="seconds"
        )
    )

    registry = _jsonl_load(
        registry_path
    )

    previous_entry = next(
        (
            item
            for item in registry
            if item.get(
                "document_id"
            )
            == document_id
        ),
        None,
    )

    # --------------------------------------------------------------
    # Duplicado verdadero:
    # ya procesado con el parser actual.
    # --------------------------------------------------------------
    if (
        previous_entry
        and previous_entry.get(
            "file_hash"
        )
        == file_hash
        and str(
            previous_entry.get(
                "parser_version",
                "",
            )
        )
        == PARSER_VERSION
    ):
        return {
            "ok": False,
            "message": (
                "Este archivo ya fue procesado "
                f"con el parser {PARSER_VERSION} "
                "y no requiere reingestión."
            ),
            "chunk_count": 0,
            "reprocessed": False,
        }

    # --------------------------------------------------------------
    # Determinar si es reprocesamiento.
    # --------------------------------------------------------------
    reprocessed = bool(
        previous_entry
    )

    previous_parser = (
        str(
            previous_entry.get(
                "parser_version",
                "anterior/no registrado",
            )
        )
        if previous_entry
        else None
    )

    # --------------------------------------------------------------
    # Persistir original.
    # --------------------------------------------------------------
    safe_name = re.sub(
        r"[^A-Za-z0-9._-]+",
        "_",
        filename,
    )

    stored_path = (
        files_dir
        / f"{document_id}__{safe_name}"
    )

    stored_path.write_bytes(
        file_bytes
    )

    # --------------------------------------------------------------
    # Crear nuevos chunks.
    # --------------------------------------------------------------
    chunks = []
    ingestion_mode = "generic_text"
    structured_count = 0
    generic_count = 0

    if extension == ".pdf":

        try:
            rows = extract_pdf_tables(
                file_bytes
            )
        except Exception:
            rows = []

        if rows:
            structured = rows_to_chunks(
                rows,
                document_id=document_id,
                title=title,
                source_file=filename,
                document_type=document_type,
                source_status=source_status,
                access_level=access_level,
                document_date=document_date,
                version=version,
                file_hash=file_hash,
                ingested_at=ingested_at,
            )

            for item in structured:
                item[
                    "parser_version"
                ] = PARSER_VERSION

            chunks.extend(
                structured
            )

            structured_count = len(
                structured
            )

            ingestion_mode = (
                "structured_table"
            )

        pages = _extract_pdf_text(
            file_bytes
        )

        generic = _build_generic_chunks(
            pages,
            document_id=document_id,
            title=title,
            source_file=filename,
            document_type=document_type,
            source_status=source_status,
            access_level=access_level,
            document_date=document_date,
            version=version,
            file_hash=file_hash,
            ingested_at=ingested_at,
            chunk_size=chunk_size,
            overlap=overlap,
        )

        if structured_count:
            for idx, chunk in enumerate(
                generic,
                start=1,
            ):
                chunk[
                    "chunk_id"
                ] = (
                    f"{document_id}-G"
                    f"{idx:04d}"
                )
                chunk[
                    "chunk_index"
                ] = idx

        generic_count = len(
            generic
        )

        chunks.extend(
            generic
        )

    elif extension == ".csv":

        rows = extract_csv_rows(
            file_bytes,
            filename,
        )

        chunks = rows_to_chunks(
            rows,
            document_id=document_id,
            title=title,
            source_file=filename,
            document_type=document_type,
            source_status=source_status,
            access_level=access_level,
            document_date=document_date,
            version=version,
            file_hash=file_hash,
            ingested_at=ingested_at,
        )

        for item in chunks:
            item[
                "parser_version"
            ] = PARSER_VERSION

        structured_count = len(
            chunks
        )

        ingestion_mode = (
            "structured_table"
        )

    elif extension in {
        ".xlsx",
        ".xls",
    }:

        rows = extract_excel_rows(
            file_bytes,
            filename,
        )

        chunks = rows_to_chunks(
            rows,
            document_id=document_id,
            title=title,
            source_file=filename,
            document_type=document_type,
            source_status=source_status,
            access_level=access_level,
            document_date=document_date,
            version=version,
            file_hash=file_hash,
            ingested_at=ingested_at,
        )

        for item in chunks:
            item[
                "parser_version"
            ] = PARSER_VERSION

        structured_count = len(
            chunks
        )

        ingestion_mode = (
            "structured_table"
        )

    elif extension == ".docx":

        text = _extract_docx_text(
            file_bytes
        )

        chunks = _build_generic_chunks(
            [
                (
                    "Documento",
                    text,
                )
            ],
            document_id=document_id,
            title=title,
            source_file=filename,
            document_type=document_type,
            source_status=source_status,
            access_level=access_level,
            document_date=document_date,
            version=version,
            file_hash=file_hash,
            ingested_at=ingested_at,
            chunk_size=chunk_size,
            overlap=overlap,
        )

        generic_count = len(
            chunks
        )

    elif extension == ".txt":

        text = _extract_txt(
            file_bytes
        )

        chunks = _build_generic_chunks(
            [
                (
                    "Documento",
                    text,
                )
            ],
            document_id=document_id,
            title=title,
            source_file=filename,
            document_type=document_type,
            source_status=source_status,
            access_level=access_level,
            document_date=document_date,
            version=version,
            file_hash=file_hash,
            ingested_at=ingested_at,
            chunk_size=chunk_size,
            overlap=overlap,
        )

        generic_count = len(
            chunks
        )

    elif extension == ".doc":
        raise RuntimeError(
            "Convierta el archivo .DOC a .DOCX "
            "antes de ingerirlo."
        )

    if not chunks:
        raise RuntimeError(
            "No se pudo extraer contenido utilizable."
        )

    # --------------------------------------------------------------
    # Reemplazar chunks anteriores del mismo ID.
    # --------------------------------------------------------------
    existing_chunks = _jsonl_load(
        chunks_path
    )

    existing_chunks = [
        item
        for item in existing_chunks
        if item.get(
            "document_id"
        )
        != document_id
    ]

    existing_chunks.extend(
        chunks
    )

    _jsonl_write(
        chunks_path,
        existing_chunks,
    )

    # --------------------------------------------------------------
    # Reemplazar entrada del registro.
    # --------------------------------------------------------------
    registry = [
        item
        for item in registry
        if item.get(
            "document_id"
        )
        != document_id
    ]

    registry.append(
        {
            "document_id": document_id,
            "title": title,
            "source_file": filename,
            "file_extension": extension,
            "document_type": document_type,
            "source_status": source_status,
            "access_level": access_level,
            "document_date": document_date,
            "version": version,
            "file_hash": file_hash,
            "ingested_at": ingested_at,
            "chunk_count": len(
                chunks
            ),
            "structured_count": (
                structured_count
            ),
            "generic_count": (
                generic_count
            ),
            "ingestion_mode": ingestion_mode,
            "parser_version": PARSER_VERSION,
        }
    )

    _jsonl_write(
        registry_path,
        registry,
    )

    if reprocessed:
        message = (
            "Documento reprocesado correctamente. "
            f"Parser anterior: {previous_parser}. "
            f"Parser actual: {PARSER_VERSION}. "
            f"Chunks estructurados: {structured_count}. "
            f"Chunks de texto: {generic_count}."
        )
    else:
        message = (
            "Documento ingerido correctamente. "
            f"Parser: {PARSER_VERSION}. "
            f"Chunks estructurados: {structured_count}. "
            f"Chunks de texto: {generic_count}."
        )

    return {
        "ok": True,
        "message": message,
        "chunk_count": len(
            chunks
        ),
        "structured_count": (
            structured_count
        ),
        "generic_count": (
            generic_count
        ),
        "ingestion_mode": ingestion_mode,
        "parser_version": PARSER_VERSION,
        "reprocessed": reprocessed,
    }
