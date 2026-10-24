"""
ingestion.py
------------
Ingestión documental local para el Asistente IA DINDES.

Formatos soportados:
- PDF
- DOCX
- DOC (Windows + Microsoft Word)
- TXT
- CSV
- XLSX
- XLS

Flujo:
archivo -> hash -> extracción -> limpieza -> chunking -> metadatos -> JSONL

No envía documentos a servicios externos.
"""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
import tempfile
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".txt",
    ".csv",
    ".xlsx",
    ".xls",
}


@dataclass
class IngestedChunk:
    chunk_id: str
    document_id: str
    title: str
    source_file: str
    document_type: str
    section: str
    chunk_index: int
    text: str
    source_status: str
    access_level: str
    document_date: str
    version: str
    file_hash: str
    ingested_at: str


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def normalize_whitespace(text: str) -> str:
    text = str(text).replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def dataframe_to_text(
    df: pd.DataFrame,
    *,
    max_rows: int | None = None,
) -> str:
    """
    Convierte un DataFrame a texto tabular legible.
    """
    if max_rows is not None:
        df = df.head(max_rows)

    if df.empty:
        return ""

    # Evita NaN y hace strings consistentes.
    clean = df.fillna("").astype(str)

    lines = []

    headers = [str(col) for col in clean.columns]
    lines.append(" | ".join(headers))

    for _, row in clean.iterrows():
        values = [
            normalize_whitespace(value)
            for value in row.tolist()
        ]
        lines.append(" | ".join(values))

    return normalize_whitespace("\n".join(lines))


def extract_pdf(path: Path) -> list[tuple[str, str]]:
    reader = PdfReader(str(path))
    sections: list[tuple[str, str]] = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = normalize_whitespace(page.extract_text() or "")
        if text:
            sections.append(
                (f"Página {page_number}", text)
            )

    return sections


def extract_docx(path: Path) -> list[tuple[str, str]]:
    document = Document(str(path))

    sections: list[tuple[str, str]] = []
    current_heading = "Documento"
    buffer: list[str] = []

    def flush():
        nonlocal buffer
        text = normalize_whitespace(
            "\n".join(buffer)
        )
        if text:
            sections.append(
                (current_heading, text)
            )
        buffer = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        style_name = (
            paragraph.style.name
            or ""
        ).lower()

        if (
            "heading" in style_name
            or "título" in style_name
            or "titulo" in style_name
        ):
            flush()
            current_heading = text
        else:
            buffer.append(text)

    flush()

    # Incorporar también contenido de tablas.
    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):
        rows = []

        for row in table.rows:
            values = [
                normalize_whitespace(
                    cell.text
                )
                for cell in row.cells
            ]
            rows.append(
                " | ".join(values)
            )

        table_text = normalize_whitespace(
            "\n".join(rows)
        )

        if table_text:
            sections.append(
                (
                    f"Tabla {table_index}",
                    table_text,
                )
            )

    return sections


def extract_doc_legacy(path: Path) -> list[tuple[str, str]]:
    """
    Extrae .DOC antiguo.

    Estrategia principal:
    Microsoft Word vía COM en Windows (pywin32).

    Fallback:
    antiword, si está instalado y disponible en PATH.
    """

    # ---------------------------------------------------------
    # Opción 1: Microsoft Word vía COM
    # ---------------------------------------------------------
    try:
        import win32com.client  # type: ignore

        word = None
        document = None

        try:
            word = win32com.client.DispatchEx(
                "Word.Application"
            )
            word.Visible = False
            word.DisplayAlerts = 0

            document = word.Documents.Open(
                str(path.resolve()),
                ReadOnly=True,
            )

            text = normalize_whitespace(
                document.Content.Text
            )

            if text:
                return [
                    ("Documento", text)
                ]

        finally:
            if document is not None:
                document.Close(False)

            if word is not None:
                word.Quit()

    except Exception:
        pass

    # ---------------------------------------------------------
    # Opción 2: antiword
    # ---------------------------------------------------------
    try:
        process = subprocess.run(
            [
                "antiword",
                str(path),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=True,
        )

        text = normalize_whitespace(
            process.stdout
        )

        if text:
            return [
                ("Documento", text)
            ]

    except Exception:
        pass

    raise RuntimeError(
        "No fue posible extraer el archivo .DOC. "
        "En Windows, instale Microsoft Word y pywin32, "
        "o convierta el archivo a DOCX antes de ingerirlo."
    )


def extract_txt(path: Path) -> list[tuple[str, str]]:
    raw = path.read_text(
        encoding="utf-8",
        errors="replace",
    )

    text = normalize_whitespace(raw)

    return (
        [("Documento", text)]
        if text
        else []
    )


def extract_csv(path: Path) -> list[tuple[str, str]]:
    """
    Intenta primero UTF-8 y luego latin-1/cp1252.
    """
    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    ]

    last_error = None

    for encoding in encodings:
        try:
            df = pd.read_csv(
                path,
                encoding=encoding,
                sep=None,
                engine="python",
            )

            text = dataframe_to_text(df)

            if text:
                return [
                    ("CSV", text)
                ]

            return []

        except Exception as exc:
            last_error = exc

    raise RuntimeError(
        f"No fue posible leer el CSV. Último error: {last_error}"
    )


def extract_excel(path: Path) -> list[tuple[str, str]]:
    """
    XLSX -> openpyxl
    XLS  -> xlrd
    """
    suffix = path.suffix.lower()

    engine = (
        "openpyxl"
        if suffix == ".xlsx"
        else "xlrd"
    )

    workbook = pd.read_excel(
        path,
        sheet_name=None,
        engine=engine,
    )

    sections: list[tuple[str, str]] = []

    for sheet_name, df in workbook.items():
        text = dataframe_to_text(df)

        if text:
            sections.append(
                (
                    f"Hoja: {sheet_name}",
                    text,
                )
            )

    return sections


def extract_sections(
    path: Path,
) -> list[tuple[str, str]]:

    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Formato no soportado: {suffix}. "
            f"Formatos permitidos: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if suffix == ".pdf":
        return extract_pdf(path)

    if suffix == ".docx":
        return extract_docx(path)

    if suffix == ".doc":
        return extract_doc_legacy(path)

    if suffix == ".txt":
        return extract_txt(path)

    if suffix == ".csv":
        return extract_csv(path)

    if suffix in {".xlsx", ".xls"}:
        return extract_excel(path)

    return []


def chunk_text(
    text: str,
    *,
    chunk_size: int = 1200,
    overlap: int = 200,
) -> list[str]:

    if chunk_size <= 0:
        raise ValueError(
            "chunk_size debe ser mayor que cero."
        )

    if (
        overlap < 0
        or overlap >= chunk_size
    ):
        raise ValueError(
            "overlap debe ser >= 0 "
            "y menor que chunk_size."
        )

    text = normalize_whitespace(text)

    if not text:
        return []

    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0

    while start < len(text):
        target_end = min(
            start + chunk_size,
            len(text),
        )

        end = target_end

        if target_end < len(text):
            search_start = max(
                start
                + int(
                    chunk_size * 0.65
                ),
                start,
            )

            candidates = [
                text.rfind(
                    "\n\n",
                    search_start,
                    target_end,
                ),
                text.rfind(
                    ". ",
                    search_start,
                    target_end,
                ),
                text.rfind(
                    "; ",
                    search_start,
                    target_end,
                ),
                text.rfind(
                    ", ",
                    search_start,
                    target_end,
                ),
                text.rfind(
                    " ",
                    search_start,
                    target_end,
                ),
            ]

            best = max(candidates)

            if best > start:
                end = best + 1

        chunk = text[
            start:end
        ].strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = max(
            end - overlap,
            start + 1,
        )

    return chunks


def build_chunks(
    path: Path,
    *,
    document_id: str,
    title: str,
    document_type: str = "Documento técnico",
    source_status: str = "vigente",
    access_level: str = "autorizado",
    document_date: str = "",
    version: str = "1.0",
    file_hash: str,
    chunk_size: int = 1200,
    overlap: int = 200,
) -> list[IngestedChunk]:

    sections = extract_sections(path)

    ingested_at = (
        datetime.now()
        .isoformat(
            timespec="seconds"
        )
    )

    result = []
    global_index = 1

    for section, text in sections:

        chunks = chunk_text(
            text,
            chunk_size=chunk_size,
            overlap=overlap,
        )

        for chunk in chunks:

            result.append(
                IngestedChunk(
                    chunk_id=(
                        f"{document_id}"
                        f"-C{global_index:04d}"
                    ),
                    document_id=document_id,
                    title=title,
                    source_file=path.name,
                    document_type=document_type,
                    section=section,
                    chunk_index=global_index,
                    text=chunk,
                    source_status=(
                        source_status.lower()
                    ),
                    access_level=(
                        access_level.lower()
                    ),
                    document_date=document_date,
                    version=version,
                    file_hash=file_hash,
                    ingested_at=ingested_at,
                )
            )

            global_index += 1

    return result


def append_jsonl(
    chunks: Iterable[IngestedChunk],
    output_path: Path,
) -> int:

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    count = 0

    with output_path.open(
        "a",
        encoding="utf-8",
    ) as file:

        for chunk in chunks:

            file.write(
                json.dumps(
                    asdict(chunk),
                    ensure_ascii=False,
                )
                + "\n"
            )

            count += 1

    return count


def load_jsonl(path: Path) -> list[dict]:

    if not path.exists():
        return []

    rows = []

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:

        for line in file:

            line = line.strip()

            if not line:
                continue

            rows.append(
                json.loads(line)
            )

    return rows


def file_hash_exists(
    registry_path: Path,
    file_hash: str,
) -> bool:

    if not registry_path.exists():
        return False

    for row in load_jsonl(
        registry_path
    ):
        if (
            row.get("file_hash")
            == file_hash
        ):
            return True

    return False


def register_document(
    registry_path: Path,
    metadata: dict,
) -> None:

    registry_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with registry_path.open(
        "a",
        encoding="utf-8",
    ) as file:

        file.write(
            json.dumps(
                metadata,
                ensure_ascii=False,
            )
            + "\n"
        )
